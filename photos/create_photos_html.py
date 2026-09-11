#!/usr/bin/env python3
"""
Create photos.html from pdf_files.xlsx.

Excel columns:
A = Year
B = Month
C = Day
D = Title
E = File Name
"""

from pathlib import Path
from urllib.parse import quote
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: openpyxl is not installed.")
    print("Install it with: pip install openpyxl")
    input("Press Enter to exit...")
    sys.exit(1)

EXCEL_FILE = "pdf_files.xlsx"
TEMPLATE_FILE = "photos_placeholder.html"
OUTPUT_FILE = "photos.html"
PDF_DIRECTORY_URL = "https://soderstrand.com/photos/"
PDFJS_VIEWER = "https://mozilla.github.io/pdf.js/web/viewer.html?file="

script_directory = Path(__file__).resolve().parent
excel_path = script_directory / EXCEL_FILE
template_path = script_directory / TEMPLATE_FILE
output_path = script_directory / OUTPUT_FILE

print("Reading Excel File")
print(f"Excel file being read: {EXCEL_FILE}")

try:
    if not excel_path.exists():
        raise FileNotFoundError(
            f"Could not find '{EXCEL_FILE}' in the same directory as this program."
        )

    workbook = load_workbook(excel_path, data_only=True)
    worksheet = workbook.active

    headers = [
        "" if cell.value is None else str(cell.value).strip()
        for cell in worksheet[1]
    ]
    expected_headers = ["Year", "Month", "Day", "Title", "File Name"]

    if headers[:5] != expected_headers:
        raise ValueError(
            "The first five Excel column headings must be: "
            "Year, Month, Day, Title, File Name"
        )

    records = []

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        if all(value is None or str(value).strip() == "" for value in row[:5]):
            continue

        title = "" if row[3] is None else str(row[3]).strip()
        filename = "" if row[4] is None else str(row[4]).strip()

        if not title:
            raise ValueError(f"Row {row_number}: Title (column D) is blank.")
        if not filename:
            raise ValueError(f"Row {row_number}: File Name (column E) is blank.")

        records.append({
            "year": row[0],
            "month": row[1],
            "day": row[2],
            "title": title,
            "filename": filename,
        })

    workbook.close()
    print(f"Records read: {len(records)}")
    print("Excel File Read Successfully")

except Exception as error:
    print(f"ERROR reading Excel file: {error}")
    input("Press Enter to exit...")
    sys.exit(1)

try:
    if not template_path.exists():
        raise FileNotFoundError(
            f"Could not find '{TEMPLATE_FILE}' in the same directory as this program."
        )
    template = template_path.read_text(encoding="utf-8")

except Exception as error:
    print(f"ERROR reading HTML template: {error}")
    input("Press Enter to exit...")
    sys.exit(1)

photo_links = []

for record in records:
    title = record["title"]
    filename = record["filename"]

    # Encode the filename for the direct PDF URL.
    encoded_filename = quote(filename, safe="")
    pdf_url = PDF_DIRECTORY_URL + encoded_filename

    # Encode the complete PDF URL as PDF.js's file parameter.
    pdfjs_url = PDFJS_VIEWER + quote(pdf_url, safe="")

    entry = (
        '        <div class="photo-entry">\n'
        f'            <h3>{title}</h3>\n'
        '            <p>\n'
        f'                <a href="{pdfjs_url}" target="_blank" rel="noopener">\n'
        '                    View PDF\n'
        '                </a>\n'
        '                &nbsp;|&nbsp;\n'
        f'                <a href="{pdf_url}" download>\n'
        '                    Download PDF\n'
        '                </a>\n'
        '            </p>\n'
        '        </div>'
    )
    photo_links.append(entry)

links_html = "\n\n".join(photo_links)

try:
    main_start = template.index("<main")
    main_open_end = template.index(">", main_start) + 1
    main_end = template.index("</main>", main_open_end)

    new_main = (
        '<main class="placeholder-page">\n'
        '    <h2>Photos</h2>\n\n'
        '    <p>\n'
        '        Select a photo collection below to open it in the PDF viewer.\n'
        '        Each PDF can also be downloaded directly.\n'
        '    </p>\n\n'
        f'{links_html}\n'
        '</main>'
    )

    html = (
        template[:main_start]
        + new_main
        + template[main_end + len("</main>"):]
    )

except ValueError:
    print("ERROR: Could not locate the <main>...</main> section in the template.")
    input("Press Enter to exit...")
    sys.exit(1)

try:
    output_path.write_text(html, encoding="utf-8")
    print("photos.html created")
except Exception as error:
    print(f"ERROR creating photos.html: {error}")
    input("Press Enter to exit...")
    sys.exit(1)

input("Press Enter to exit...")
